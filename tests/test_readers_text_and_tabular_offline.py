import pytest
from openpyxl import Workbook

from ai_engine.readers import read_csv, read_markdown, read_text, read_xlsx


def test_read_text_returns_utf8_text_and_metadata(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Olá, mundo!\nSecond line", encoding="utf-8")

    document = read_text(path)

    assert document.source_path == path
    assert document.text == "Olá, mundo!\nSecond line"
    assert document.metadata == {
        "format": "txt",
        "filename": "notes.txt",
    }


@pytest.mark.parametrize("extension", [".md", ".markdown"])
def test_read_markdown_preserves_source_text(extension, tmp_path):
    path = tmp_path / f"guide{extension}"
    path.write_text("# Title\n\n- item", encoding="utf-8")

    document = read_markdown(path)

    assert document.source_path == path
    assert document.text == "# Title\n\n- item"
    assert document.metadata == {
        "format": "markdown",
        "filename": path.name,
    }


def test_read_csv_detects_delimiter_strips_cells_and_reports_rows(tmp_path):
    path = tmp_path / "people.csv"
    path.write_text(
        "\ufeffname; city\nAlice; São Paulo\nBob; Recife\n",
        encoding="utf-8",
    )

    document = read_csv(path)

    assert len(document.tables) == 1
    assert document.tables[0].name == "people"
    assert document.tables[0].source == "people.csv"
    assert document.tables[0].rows == [
        ["name", "city"],
        ["Alice", "São Paulo"],
        ["Bob", "Recife"],
    ]
    assert document.metadata == {
        "format": "csv",
        "filename": "people.csv",
        "row_count": 3,
    }


def test_read_xlsx_ignores_empty_rows_and_reports_worksheet_metadata(tmp_path):
    path = tmp_path / "workbook.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["name", "value"])
    worksheet.append([None, None])
    worksheet.append(["alpha", 10])
    workbook.save(path)

    document = read_xlsx(path)

    assert len(document.tables) == 1
    assert document.tables[0].rows == [
        ["name", "value"],
        ["alpha", "10"],
    ]
    assert document.metadata == {
        "format": "xlsx",
        "filename": "workbook.xlsx",
        "sheet_count": 1,
        "sheet_names": ["Data"],
        "sheet_row_counts": {"Data": 2},
        "sheet_column_counts": {"Data": 2},
    }


@pytest.mark.parametrize(
    ("reader", "filename"),
    [
        (read_text, "missing.txt"),
        (read_markdown, "missing.md"),
        (read_csv, "missing.csv"),
        (read_xlsx, "missing.xlsx"),
    ],
)
def test_text_and_tabular_readers_reject_missing_files(reader, filename, tmp_path):
    with pytest.raises(FileNotFoundError, match="File not found"):
        reader(tmp_path / filename)


@pytest.mark.parametrize(
    ("reader", "filename", "message"),
    [
        (read_markdown, "notes.txt", "Expected a Markdown file"),
        (read_csv, "table.txt", "Expected a .csv file"),
        (read_xlsx, "workbook.txt", "Expected an .xlsx or .xlsm file"),
    ],
)
def test_text_and_tabular_readers_reject_wrong_extensions(
    reader,
    filename,
    message,
    tmp_path,
):
    path = tmp_path / filename
    path.write_text("content", encoding="utf-8")

    with pytest.raises(ValueError, match=message):
        reader(path)

